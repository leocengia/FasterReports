"""
Genera wow_CaseType_Deepdive.xlsx - analisi mensile WoW per case type e canale.

Input:  raw_data/<filename>.csv  (separatore ;, decimale ,)
        Colonne attese: Case Type, Case AHT (mins), Distinct Cases,
                        Case Origin (group) o Case Channel,
                        MonthKey (opzionale, es. 2026-06-01),
                        OOT_flag (opzionale, derivato automaticamente se assente)
Output: output/wow_CaseType_Deepdive_<periodo>.xlsx

Uso:
    python src/wow_casetype_deepdive.py "raw_data/dataset.csv"
    python src/wow_casetype_deepdive.py "raw_data/dataset.csv" --month 2026-06
    python src/wow_casetype_deepdive.py "raw_data/dataset.csv" --month 2026-06 --target-phone 19.98
"""

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── Parametri configurabili ───────────────────────────────────────────────────
DEFAULT_TGT_PHONE = 19.98
DEFAULT_TGT_NL    = 18.96
MIN_VOLUME        = 20      # volume minimo per classificazione
HIGHLY_THRESHOLD  = 0.75   # soglia score → Highly Actionable
ACTION_THRESHOLD  = 0.50   # soglia score → Actionable

NON_LIVE_CHANNELS = {"Contact Us", "Market Management", "Others", "Email"}
PHONE_CHANNELS    = {"Phone"}

# ── Colori (dal template) ─────────────────────────────────────────────────────
C_HDR_BG      = "000099"   # blu scuro intestazioni
C_HDR_FG      = "FFFFFF"   # bianco testo intestazioni
C_PARAM_BG    = "4472C4"   # blu medio celle parametro
C_HELPER_BG   = "5B9BD5"   # blu chiaro helper
C_NODATA_BG   = "EAEAEA"   # grigio dati assenti
C_HA_PHONE_BG = "CB2C30"   # rosso scuro Highly Actionable (Phone)
C_HA_PHONE_FG = "FFFFFF"
C_ACT_NL_BG   = "FFC60B"   # giallo Actionable (Non-live)
C_ACT_NL_FG   = "000000"
C_ACT_PH_BG   = "D9E1F2"   # azzurro chiaro Actionable (Phone)
C_ACT_PH_FG   = "000000"
C_ACT_NL2_BG  = "FFF2CC"   # giallo chiaro Actionable (Non-live, alternativa)
C_PROC_BG     = None        # Process Driven: nessun fill

THIN   = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ── Helpers ───────────────────────────────────────────────────────────────────

def detect_col(df, *candidates):
    for name in candidates:
        if name in df.columns:
            return name
    return None


def _cell(ws, row, col, value=None, *, bold=False, size=11,
          align="center", fmt=None, bg=None, fg="000000", wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(bold=bold, name="Calibri", size=size, color=fg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border    = BORDER
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if fmt:
        c.number_format = fmt
    return c


def _hdr(ws, row, col, value, **kwargs):
    return _cell(ws, row, col, value, bold=True, bg=C_HDR_BG, fg=C_HDR_FG, **kwargs)


# ── Calcolo statistiche per case type ─────────────────────────────────────────

def compute_stats(df, aht_col, cases_col, type_col, oot_col=None, target=None):
    """
    Restituisce DataFrame con: case_type, aht, volume, impact, oot_rate,
    std_dev, p90p10  per ogni case type presente nei dati.
    """
    rows = []
    total_cases = df[cases_col].sum()

    for ct, grp in df.groupby(type_col):
        vol = grp[cases_col].sum()
        if vol == 0:
            continue
        avg_aht = (grp[aht_col] * grp[cases_col]).sum() / vol

        # std_dev: calcolata sulle righe (approssimazione per dati aggregati)
        std_dev = grp[aht_col].std() if len(grp) > 1 else 0.0
        p90     = grp[aht_col].quantile(0.9)
        p10     = grp[aht_col].quantile(0.1)
        p90p10  = p90 - p10

        # OOT rate
        if oot_col and oot_col in grp.columns:
            oot_rate = (grp[oot_col] * grp[cases_col]).sum() / vol
        elif target is not None:
            # approssimazione: % righe con AHT > target, ponderata per volume
            oot_mask = grp[aht_col] > target
            oot_rate = (grp.loc[oot_mask, cases_col].sum() / vol)
        else:
            oot_rate = 0.0

        impact = (avg_aht * vol) / (df[aht_col] * df[cases_col]).sum() \
                 if (df[aht_col] * df[cases_col]).sum() > 0 else 0.0

        rows.append({
            "case_type": ct,
            "aht":       avg_aht,
            "volume":    vol,
            "impact":    impact,
            "oot_rate":  oot_rate,
            "std_dev":   std_dev if not np.isnan(std_dev) else 0.0,
            "p90p10":    p90p10  if not np.isnan(p90p10)  else 0.0,
        })

    if not rows:
        return pd.DataFrame(columns=["case_type", "aht", "volume", "impact",
                                     "oot_rate", "std_dev", "p90p10"])
    return pd.DataFrame(rows).sort_values("case_type").reset_index(drop=True)


# ── Algoritmo di classificazione (da tblClass del template) ──────────────────

def classify(stats_df, target_aht, min_vol=MIN_VOLUME,
             highly_thr=HIGHLY_THRESHOLD, act_thr=ACTION_THRESHOLD):
    """
    Calcola score e categoria per ogni case type.
    Score = 0.25*impact_rank + 0.2*opp_rank + 0.15*vol_rank
             + 0.2*oot_rank + 0.2*variability_score
    Variability = 0.4*std_rank + 0.6*p90p10_rank
    """
    df = stats_df.copy()
    df["gap"]  = (df["aht"] - target_aht).clip(lower=0)
    df["opp"]  = df["gap"] * df["volume"]

    def pct_rank(s):
        return s.rank(pct=True, method="average")

    df["impact_rank"]  = pct_rank(df["impact"])
    df["opp_rank"]     = pct_rank(df["opp"])
    df["vol_rank"]     = pct_rank(df["volume"])
    df["oot_rank"]     = pct_rank(df["oot_rate"])
    df["std_rank"]     = pct_rank(df["std_dev"])
    df["p90p10_rank"]  = pct_rank(df["p90p10"])
    df["var_score"]    = 0.4 * df["std_rank"] + 0.6 * df["p90p10_rank"]
    df["score"]        = (0.25 * df["impact_rank"]
                          + 0.20 * df["opp_rank"]
                          + 0.15 * df["vol_rank"]
                          + 0.20 * df["oot_rank"]
                          + 0.20 * df["var_score"])

    def cat(row):
        if row["volume"] < min_vol:
            return "Process Driven"
        if row["score"] >= highly_thr:
            return "Highly Actionable"
        if row["score"] >= act_thr:
            return "Actionable"
        return "Process Driven"

    df["categoria"] = df.apply(cat, axis=1)
    return df


# ── Scrittura foglio DATASET ───────────────────────────────────────────────────

def write_dataset_sheet(ws, df):
    hdr_font = Font(name="Calibri", size=11)
    for c_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=c_idx, value=col_name).font = hdr_font
    for r_idx, row in enumerate(df.itertuples(index=False), 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)


# ── Scrittura foglio Monthly ───────────────────────────────────────────────────

def write_monthly_sheet(ws, stats_phone, stats_nl, period_label):
    """
    Layout: sezione Phone (righe 4-36+) poi sezione Non-live (righe 39-71+).
    """
    col_widths = {
        "A": 33.29, "B": 10.0, "C": 7.71, "D": 10.43,
        "E": 16.71, "F": 9.14, "G": 8.57,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    headers = ["aht (mins)", "volume", "impact (%)", "out_of_target (%)", "std_dev", "p90-p10"]

    def write_section(start_row, label, stats):
        # Titolo sezione
        ws.merge_cells(start_row=start_row, start_column=1,
                       end_row=start_row,   end_column=1)
        _hdr(ws, start_row, 1, label, align="center")
        ws.merge_cells(start_row=start_row, start_column=2,
                       end_row=start_row,   end_column=7)
        _hdr(ws, start_row, 2, period_label, align="center")

        # Intestazioni colonne
        hdr_row = start_row + 1
        _hdr(ws, hdr_row, 1, "case_type")
        for i, h in enumerate(headers):
            _hdr(ws, hdr_row, 2 + i, h)

        # Dati
        for row_i, row in stats.iterrows():
            xl_row = hdr_row + 1 + row_i
            _cell(ws, xl_row, 1, row["case_type"], bold=True, align="left")
            if row["volume"] == 0:
                for c in range(2, 8):
                    _cell(ws, xl_row, c, "N/A", bold=True, bg=C_NODATA_BG)
            else:
                _cell(ws, xl_row, 2, row["aht"],      fmt="0.00")
                _cell(ws, xl_row, 3, row["volume"],   fmt="0")
                _cell(ws, xl_row, 4, row["impact"],   fmt="0.00%")
                _cell(ws, xl_row, 5, row["oot_rate"], fmt="0.00%")
                _cell(ws, xl_row, 6, row["std_dev"],  fmt="0.00")
                _cell(ws, xl_row, 7, row["p90p10"],   fmt="0.00")
        return hdr_row + 1 + len(stats)

    end_phone = write_section(4, "[ Phone ]", stats_phone)
    write_section(end_phone + 2, "[ Non-live ]", stats_nl)


# ── Scrittura foglio manualEXPORT ─────────────────────────────────────────────

def write_manual_export(ws, stats_phone, stats_nl, period_label,
                         tgt_phone, tgt_nl):
    """
    Layout (dal template):
      B     : case_type (header verticale B3:B4)
      C-H   : Phone (header C3:H3 = period + " [ Phone ]")
      I     : spacer
      J-O   : Non-live (header J3:O3)
      P     : spacer
      Q     : case_type top priorities
      R-W   : metriche top priorities Phone
    """
    col_widths = {
        "A": 5.14, "B": 32.0, "C": 10.0, "D": 7.71, "E": 10.43,
        "F": 16.71, "G": 8.0, "H": 8.0, "I": 1.29,
        "J": 10.0, "K": 7.71, "L": 10.43, "M": 16.71, "N": 8.0, "O": 8.0,
        "P": 3.86, "Q": 27.43, "R": 10.0, "S": 7.71, "T": 10.43,
        "U": 16.71, "V": 8.0, "W": 8.0,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # Altezze righe
    ws.row_dimensions[3].height = 15.75
    ws.row_dimensions[4].height = 15.75

    metric_headers = ["aht (mins)", "volume", "impact (%)", "out_of_target (%)",
                      "std_dev", "p90-p10"]

    # ── Riga 3: intestazioni sezioni ──────────────────────────────────────────
    ws.merge_cells("B3:B4")
    _hdr(ws, 3, 2, "case_type", wrap=True)

    ws.merge_cells("C3:H3")
    _hdr(ws, 3, 3, f"{period_label} - [ Phone ]")

    ws.merge_cells("J3:O3")
    _hdr(ws, 3, 10, f"{period_label} - [ Non-live ]")

    ws.merge_cells("R3:W3")
    _hdr(ws, 3, 18, f"[ Phone ] — Top Priority")

    # ── Riga 4: intestazioni metriche ─────────────────────────────────────────
    for i, h in enumerate(metric_headers):
        _hdr(ws, 4, 3  + i, h)   # Phone C4:H4
        _hdr(ws, 4, 10 + i, h)   # Non-live J4:O4
        _hdr(ws, 4, 18 + i, h)   # Top priority R4:W4

    # Intestazione case_type top priority
    _hdr(ws, 4, 17, "case_type")

    # ── Unisce tutti i case types (Phone + Non-live) ───────────────────────────
    all_types = sorted(
        set(stats_phone["case_type"]) | set(stats_nl["case_type"])
    )
    ph_by_type = stats_phone.set_index("case_type")
    nl_by_type = stats_nl.set_index("case_type")

    def pick_colors(cat, channel):
        if cat == "Highly Actionable":
            return (C_HA_PHONE_BG, C_HA_PHONE_FG)
        if cat == "Actionable":
            if channel == "Phone":
                return (C_ACT_PH_BG, C_ACT_PH_FG)
            else:
                return (C_ACT_NL_BG, C_ACT_NL_FG)
        return (None, "000000")

    # ── Righe dati (da riga 5) ─────────────────────────────────────────────────
    for row_i, ct in enumerate(all_types):
        xl_row = 5 + row_i

        # Colonna B: case_type
        _cell(ws, xl_row, 2, ct, bold=True, align="center")

        # Sezione Phone (C-H)
        if ct in ph_by_type.index:
            r = ph_by_type.loc[ct]
            cat  = r.get("categoria", "Process Driven")
            bg, fg = pick_colors(cat, "Phone")
            if r["volume"] == 0:
                for c in range(3, 9):
                    _cell(ws, xl_row, c, "N/A", bold=True, bg=C_NODATA_BG)
            else:
                vals = [r["aht"], r["volume"], r["impact"], r["oot_rate"],
                        r["std_dev"], r["p90p10"]]
                fmts = ["0.00", "0", "0.00%", "0.00%", "0.00", "0.00"]
                for c_i, (val, fmt) in enumerate(zip(vals, fmts)):
                    _cell(ws, xl_row, 3 + c_i, val, bold=(bg is not None),
                          bg=bg, fg=fg, fmt=fmt)
        else:
            for c in range(3, 9):
                _cell(ws, xl_row, c, "N/A", bold=True, bg=C_NODATA_BG)

        # Sezione Non-live (J-O)
        if ct in nl_by_type.index:
            r = nl_by_type.loc[ct]
            cat  = r.get("categoria", "Process Driven")
            bg, fg = pick_colors(cat, "Non-live")
            if r["volume"] == 0:
                for c in range(10, 16):
                    _cell(ws, xl_row, c, "N/A", bold=True, bg=C_NODATA_BG)
            else:
                vals = [r["aht"], r["volume"], r["impact"], r["oot_rate"],
                        r["std_dev"], r["p90p10"]]
                fmts = ["0.00", "0", "0.00%", "0.00%", "0.00", "0.00"]
                for c_i, (val, fmt) in enumerate(zip(vals, fmts)):
                    _cell(ws, xl_row, 10 + c_i, val, bold=(bg is not None),
                          bg=bg, fg=fg, fmt=fmt)
        else:
            for c in range(10, 16):
                _cell(ws, xl_row, c, "N/A", bold=True, bg=C_NODATA_BG)

    # ── Sezione Top Priority Phone (Q-W) ──────────────────────────────────────
    top_phone = (stats_phone[stats_phone["categoria"].isin(
                     ["Highly Actionable", "Actionable"])]
                 .sort_values("score", ascending=False)
                 .head(10)
                 .reset_index(drop=True))

    for row_i, r in top_phone.iterrows():
        xl_row = 5 + row_i
        cat = r.get("categoria", "Process Driven")
        bg, fg = pick_colors(cat, "Phone")
        _cell(ws, xl_row, 17, r["case_type"], bold=True, align="left",
              bg=bg, fg=fg)
        vals = [r["aht"], r["volume"], r["impact"], r["oot_rate"],
                r["std_dev"], r["p90p10"]]
        fmts = ["0.00", "0", "0.00%", "0.00%", "0.00", "0.00"]
        for c_i, (val, fmt) in enumerate(zip(vals, fmts)):
            _cell(ws, xl_row, 18 + c_i, val, bold=False,
                  bg=bg if bg else None, fg=fg, fmt=fmt)


# ── Scrittura foglio tblClass (scoring dettagliato) ───────────────────────────

def write_tblclass_sheet(ws, classified_phone, classified_nl):
    combined = pd.concat([
        classified_phone.assign(channel="Phone"),
        classified_nl.assign(channel="Non-live"),
    ]).reset_index(drop=True)

    cols = ["channel", "case_type", "aht", "volume", "impact", "oot_rate",
            "std_dev", "p90p10", "gap", "opp", "score", "categoria"]
    headers_pretty = ["Channel", "Case Type", "AHT", "Volume", "Impact",
                      "OOT Rate", "Std Dev", "P90-P10", "Gap vs Target",
                      "Opportunity Mins", "Score", "Categoria"]

    hdr_font = Font(name="Calibri", size=11, bold=True, color=C_HDR_FG)
    hdr_fill = PatternFill("solid", fgColor=C_HDR_BG)
    hdr_align = Alignment(horizontal="center", vertical="center")

    for c_idx, h in enumerate(headers_pretty, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = BORDER

    ws.column_dimensions["A"].width = 11.0
    ws.column_dimensions["B"].width = 32.0
    for col_letter in "CDEFGHIJKL":
        ws.column_dimensions[col_letter].width = 13.0

    fmts = [None, None, "0.00", "0", "0.00%", "0.00%",
            "0.00", "0.00", "0.00", "0.00", "0.00", None]

    for r_idx, row in combined[cols].iterrows():
        xl_row = 2 + r_idx
        for c_idx, (val, fmt) in enumerate(zip(row, fmts), 1):
            cell = ws.cell(row=xl_row, column=c_idx, value=val)
            cell.font      = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal="center" if c_idx != 2 else "left",
                                       vertical="center")
            cell.border    = BORDER
            if fmt:
                cell.number_format = fmt


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Genera wow_CaseType_Deepdive.xlsx")
    parser.add_argument("input",
                        help="Path al CSV raw (export completo casi)")
    parser.add_argument("--month", default=None,
                        help="Filtra per mese (es. 2026-06). Vuoto = tutti i dati.")
    parser.add_argument("--target-phone", type=float, default=DEFAULT_TGT_PHONE,
                        help=f"Target AHT Phone (default {DEFAULT_TGT_PHONE})")
    parser.add_argument("--target-nonlive", type=float, default=DEFAULT_TGT_NL,
                        help=f"Target AHT Non-live (default {DEFAULT_TGT_NL})")
    parser.add_argument("--min-volume", type=int, default=MIN_VOLUME,
                        help=f"Volume minimo per classificazione (default {MIN_VOLUME})")
    args = parser.parse_args()

    csv_path = Path(args.input)
    if not csv_path.exists():
        print(f"Errore: file non trovato: {csv_path}", file=sys.stderr)
        sys.exit(1)

    print("Lettura dati...")
    df = pd.read_csv(csv_path, sep=";", encoding="utf-8-sig",
                     decimal=",", low_memory=False)

    # Rileva colonne
    aht_col  = detect_col(df, "Case AHT (mins)", "DS", "aht")
    cas_col  = detect_col(df, "Distinct Cases",  "DX", "cases")
    typ_col  = detect_col(df, "Case Type",       "Z",  "case_type")
    ch_col   = detect_col(df, "Case Origin (group)", "U", "Case Channel", "S")
    mth_col  = detect_col(df, "MonthKey", "EI", "Date Viewpoint", "A")
    oot_col  = detect_col(df, "OOT_flag", "EJ")

    missing = [n for n, c in [("Case AHT",      aht_col),
                               ("Distinct Cases", cas_col),
                               ("Case Type",     typ_col)] if not c]
    if missing:
        print(f"Errore: colonne obbligatorie non trovate: {missing}", file=sys.stderr)
        print(f"Colonne disponibili: {list(df.columns[:20])}...", file=sys.stderr)
        sys.exit(1)

    df[aht_col] = pd.to_numeric(df[aht_col], errors="coerce").fillna(0)
    df[cas_col] = pd.to_numeric(df[cas_col], errors="coerce").fillna(0)
    df = df[df[cas_col] > 0].copy()

    # Mappatura canale
    if ch_col:
        if ch_col in ("Case Channel", "S"):
            def map_ch(ch):
                if ch in PHONE_CHANNELS:    return "Phone"
                if ch in NON_LIVE_CHANNELS: return "Non-live"
                return "Other"
            df["_channel"] = df[ch_col].apply(map_ch)
            ch_col = "_channel"

    # Filtro per mese
    period_label = "All"
    if args.month and mth_col:
        df[mth_col] = df[mth_col].astype(str)
        df = df[df[mth_col].str.startswith(args.month)]
        period_label = args.month
        print(f"Filtro mese: {args.month} → {len(df)} righe")
    elif args.month:
        print("Avviso: colonna MonthKey non trovata, uso tutti i dati.", file=sys.stderr)

    # Dividi per canale
    if ch_col:
        df_phone = df[df[ch_col] == "Phone"].copy()
        df_nl    = df[df[ch_col] == "Non-live"].copy()
    else:
        print("Avviso: colonna canale non trovata, uso tutti i dati per entrambi.",
              file=sys.stderr)
        df_phone = df.copy()
        df_nl    = df.copy()

    # Statistiche
    print("Calcolo statistiche...")
    stats_ph = compute_stats(df_phone, aht_col, cas_col, typ_col,
                              oot_col, args.target_phone)
    stats_nl = compute_stats(df_nl,    aht_col, cas_col, typ_col,
                              oot_col, args.target_nonlive)

    # Classificazione
    print("Classificazione case types...")
    cl_ph = classify(stats_ph, args.target_phone,  args.min_volume)
    cl_nl = classify(stats_nl, args.target_nonlive, args.min_volume)

    # Excel output
    wb = openpyxl.Workbook()

    ws_data = wb.active
    ws_data.title = "DATASET"
    print("Scrittura DATASET...")
    write_dataset_sheet(ws_data, df)

    ws_monthly = wb.create_sheet("Monthly")
    print("Scrittura Monthly...")
    write_monthly_sheet(ws_monthly, stats_ph, stats_nl, period_label)

    ws_class = wb.create_sheet("tblClass")
    print("Scrittura tblClass...")
    write_tblclass_sheet(ws_class, cl_ph, cl_nl)

    ws_export = wb.create_sheet("manualEXPORT")
    print("Scrittura manualEXPORT...")
    write_manual_export(ws_export, cl_ph, cl_nl, period_label,
                        args.target_phone, args.target_nonlive)

    wb.active = ws_export

    safe_period = period_label.replace(" ", "_").replace("/", "-")
    out_dir  = Path("output")
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / f"wow_CaseType_Deepdive_{safe_period}.xlsx"
    wb.save(out_path)

    print(f"\nSalvato: {out_path}")
    print(f"\nRiepilogo:")
    print(f"  Righe raw data:    {len(df)}")
    print(f"  Case Types Phone:  {len(stats_ph)}")
    print(f"  Case Types NL:     {len(stats_nl)}")

    ha_ph = (cl_ph["categoria"] == "Highly Actionable").sum()
    ac_ph = (cl_ph["categoria"] == "Actionable").sum()
    ha_nl = (cl_nl["categoria"] == "Highly Actionable").sum()
    ac_nl = (cl_nl["categoria"] == "Actionable").sum()
    print(f"\n  Phone   → Highly Actionable: {ha_ph}  |  Actionable: {ac_ph}")
    print(f"  Non-live → Highly Actionable: {ha_nl}  |  Actionable: {ac_nl}")


if __name__ == "__main__":
    main()
